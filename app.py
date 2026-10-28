from contextlib import contextmanager
from flask import Flask, render_template, request, redirect, url_for, send_file, flash, session, abort, jsonify
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from flask_wtf.csrf import CSRFProtect
import sqlite3
import qrcode
import qrcode.image.svg
import os
import secrets
from werkzeug.security import generate_password_hash, check_password_hash
import re
import unicodedata
from io import BytesIO
from pathlib import Path

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config["DATABASE_PATH"] = Path(__file__).resolve().with_name("database.db")
app.config["PUBLIC_BASE_URL"] = os.environ.get("PUBLIC_BASE_URL", "").strip().rstrip("/")
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = app.config["PUBLIC_BASE_URL"].startswith("https://")
app.config["RATELIMIT_STORAGE_URI"] = os.environ.get(
    "RATELIMIT_STORAGE_URI", "memory://"
)

PORTA = 5000
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024
CSRFProtect(app)
limiter = Limiter(
    get_remote_address,
    app=app,
    storage_uri=app.config["RATELIMIT_STORAGE_URI"],
)


@app.errorhandler(413)
def arquivo_acima_limite(_erro):
    sala_id = (request.view_args or {}).get("sala_id")
    if sala_id:
        flash("O arquivo deve ter no máximo 8 MB.")
        return redirect(url_for("abrir_sala", sala_id=sala_id))
    return "O arquivo deve ter no máximo 8 MB.", 413


@app.errorhandler(429)
def muitas_tentativas(_erro):
    return "Muitas tentativas. Aguarde um pouco e tente novamente.", 429


@contextmanager
def conectar_banco():
    conexao = sqlite3.connect(app.config["DATABASE_PATH"])
    conexao.row_factory = sqlite3.Row
    conexao.execute("PRAGMA foreign_keys = ON")
    try:
        with conexao:
            yield conexao
    finally:
        conexao.close()


def registrar_acao(conexao, acao, detalhe):
    conexao.execute("""
        INSERT INTO historico_acoes (usuario_id, usuario, acao, detalhe)
        VALUES (?, ?, ?, ?)
    """, (
        session.get("usuario_id"),
        session.get("nome_usuario", "Usuário"),
        acao,
        detalhe
    ))


def cor_texto_botao(cor):
    valores = [int(cor[indice:indice + 2], 16) / 255 for indice in (1, 3, 5)]
    canais = [valor / 12.92 if valor <= 0.04045 else ((valor + 0.055) / 1.055) ** 2.4 for valor in valores]
    luminancia = 0.2126 * canais[0] + 0.7152 * canais[1] + 0.0722 * canais[2]
    return "#1d1d1f" if luminancia > 0.42 else "#ffffff"

def criar_banco():
    with conectar_banco() as conexao:
        conexao.execute("""
            CREATE TABLE IF NOT EXISTS salas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome TEXT NOT NULL UNIQUE
            )
        """)

        colunas_salas = {
            linha["name"]
            for linha in conexao.execute("PRAGMA table_info(salas)")
        }

        if "numero" not in colunas_salas:
            conexao.execute("ALTER TABLE salas ADD COLUMN numero TEXT")

        if "token_publico" not in colunas_salas:
            conexao.execute("ALTER TABLE salas ADD COLUMN token_publico TEXT")

        salas_sem_token = conexao.execute("""
            SELECT id FROM salas WHERE token_publico IS NULL OR token_publico = ''
        """).fetchall()
        for sala_sem_token in salas_sem_token:
            conexao.execute(
                "UPDATE salas SET token_publico = ? WHERE id = ?",
                (secrets.token_urlsafe(24), sala_sem_token["id"])
            )
        conexao.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS idx_salas_token_publico
            ON salas(token_publico)
        """)

        conexao.execute("""
            UPDATE salas
            SET numero = CAST(id AS TEXT)
            WHERE numero IS NULL OR TRIM(numero) = ''
        """)

        conexao.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS idx_salas_numero
            ON salas(numero)
        """)

        conexao.execute("""
            CREATE TABLE IF NOT EXISTS bens (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                patrimonio TEXT NOT NULL UNIQUE,
                descricao TEXT NOT NULL,
                sala_id INTEGER NOT NULL,
                FOREIGN KEY (sala_id) REFERENCES salas(id)
            )
        """)

        colunas_bens = {
            linha["name"]
            for linha in conexao.execute("PRAGMA table_info(bens)")
        }

        if "etiquetado" not in colunas_bens:
            conexao.execute("""
                ALTER TABLE bens
                ADD COLUMN etiquetado INTEGER NOT NULL DEFAULT 0
            """)

        if "tipo" not in colunas_bens:
            conexao.execute("""
                ALTER TABLE bens
                ADD COLUMN tipo TEXT NOT NULL DEFAULT ''
            """)

        if "conferido" not in colunas_bens:
            conexao.execute("""
                ALTER TABLE bens
                ADD COLUMN conferido INTEGER NOT NULL DEFAULT 0
            """)

        if "estado" not in colunas_bens:
            conexao.execute("""
                ALTER TABLE bens
                ADD COLUMN estado TEXT NOT NULL DEFAULT ''
            """)

        conexao.execute("""
            CREATE TABLE IF NOT EXISTS tipos_bens (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome TEXT NOT NULL UNIQUE
            )
        """)

        conexao.execute("""
            CREATE TABLE IF NOT EXISTS descricoes_bens (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome TEXT NOT NULL COLLATE NOCASE UNIQUE
            )
        """)

        tipos_antigos = conexao.execute("""
            SELECT DISTINCT UPPER(TRIM(tipo)) AS nome
            FROM bens
            WHERE tipo IS NOT NULL AND TRIM(tipo) <> ''
        """).fetchall()

        for tipo in tipos_antigos:
            conexao.execute("""
                INSERT OR IGNORE INTO tipos_bens (nome)
                VALUES (?)
            """, (tipo["nome"],))

        descricoes_antigas = conexao.execute("""
            SELECT DISTINCT TRIM(descricao) AS nome
            FROM bens
            WHERE descricao IS NOT NULL AND TRIM(descricao) <> ''
        """).fetchall()

        for descricao in descricoes_antigas:
            conexao.execute("""
                INSERT OR IGNORE INTO descricoes_bens (nome)
                VALUES (?)
            """, (descricao["nome"],))

        conexao.execute("""
            CREATE TABLE IF NOT EXISTS usuarios (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome_usuario TEXT NOT NULL COLLATE NOCASE UNIQUE,
                senha_hash TEXT NOT NULL,
                perfil TEXT NOT NULL CHECK (perfil IN ('admin', 'user')),
                ativo INTEGER NOT NULL DEFAULT 1,
                criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                tema TEXT NOT NULL DEFAULT 'system' CHECK (tema IN ('system', 'light', 'dark')),
                cor_botao TEXT NOT NULL DEFAULT '#007aff',
                senha_deve_alterar INTEGER NOT NULL DEFAULT 0
            )
        """)

        conexao.execute("""
            CREATE TABLE IF NOT EXISTS historico_acoes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                usuario_id INTEGER,
                usuario TEXT NOT NULL,
                acao TEXT NOT NULL,
                detalhe TEXT NOT NULL,
                criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conexao.execute("""
            CREATE INDEX IF NOT EXISTS idx_historico_acoes_criado_em
            ON historico_acoes(criado_em DESC, id DESC)
        """)

        colunas_usuarios = {
            linha["name"] for linha in conexao.execute("PRAGMA table_info(usuarios)")
        }
        if "tema" not in colunas_usuarios:
            conexao.execute("ALTER TABLE usuarios ADD COLUMN tema TEXT NOT NULL DEFAULT 'system'")
        if "cor_botao" not in colunas_usuarios:
            conexao.execute("ALTER TABLE usuarios ADD COLUMN cor_botao TEXT NOT NULL DEFAULT '#007aff'")
        if "senha_deve_alterar" not in colunas_usuarios:
            conexao.execute(
                "ALTER TABLE usuarios ADD COLUMN senha_deve_alterar INTEGER NOT NULL DEFAULT 0"
            )
        total_usuarios = conexao.execute(
            "SELECT COUNT(*) AS total FROM usuarios"
        ).fetchone()["total"]
        if total_usuarios == 0:
            nome_admin = os.environ.get("ADMIN_USERNAME", "admin").strip()
            senha_admin = os.environ.get("ADMIN_PASSWORD", "")
            senha_gerada = not senha_admin
            if senha_gerada:
                senha_admin = secrets.token_urlsafe(24)
            if len(nome_admin) < 3 or len(nome_admin) > 50 or not re.fullmatch(r"[A-Za-z0-9._-]+", nome_admin):
                raise RuntimeError("ADMIN_USERNAME deve ter de 3 a 50 caracteres: letras, números, ponto, hífen ou sublinhado.")
            if len(senha_admin) < 12:
                raise RuntimeError("ADMIN_PASSWORD deve ter ao menos 12 caracteres.")
            conexao.execute("""
                INSERT INTO usuarios (nome_usuario, senha_hash, perfil)
                VALUES (?, ?, 'admin')
            """, (nome_admin, generate_password_hash(senha_admin)))
            if senha_gerada:
                print(f"Primeiro administrador: {nome_admin} | Senha inicial: {senha_admin}")


@app.before_request
def exigir_login():
    if request.endpoint in {
        "login", "criar_conta", "recuperar_senha", "espelho_sala", "static"
    }:
        return None

    usuario_id = session.get("usuario_id")
    if not usuario_id:
        return redirect(url_for("login"))

    with conectar_banco() as conexao:
        usuario = conexao.execute("""
            SELECT id, nome_usuario, perfil, ativo, tema, cor_botao,
                   senha_deve_alterar
            FROM usuarios WHERE id = ?
        """, (usuario_id,)).fetchone()

    if usuario is None or not usuario["ativo"]:
        session.clear()
        return redirect(url_for("login"))

    session["nome_usuario"] = usuario["nome_usuario"]
    session["perfil"] = usuario["perfil"]
    session["senha_deve_alterar"] = bool(usuario["senha_deve_alterar"])
    session["tema_usuario"] = usuario["tema"] if usuario["tema"] in {"system", "light", "dark"} else "system"
    cor_botao = usuario["cor_botao"] or "#007aff"
    if not re.fullmatch(r"#[0-9A-Fa-f]{6}", cor_botao):
        cor_botao = "#007aff"
    session["cor_botao"] = cor_botao.lower()
    session["cor_botao_texto"] = cor_texto_botao(cor_botao)
    if session["senha_deve_alterar"] and request.endpoint not in {
        "alterar_senha_propria", "sair", "static"
    }:
        return redirect(url_for("alterar_senha_propria"))
    if request.endpoint in {
        "usuarios", "criar_usuario", "alterar_perfil_usuario",
        "excluir_usuario", "redefinir_senha_usuario"
    } and usuario["perfil"] != "admin":
        abort(403)


def validar_credenciais_usuario(nome_usuario, senha):
    if not re.fullmatch(r"[A-Za-z0-9._-]{3,50}", nome_usuario):
        return "O usuário deve ter de 3 a 50 caracteres: letras, números, ponto, hífen ou sublinhado."
    if len(senha) < 12:
        return "A senha deve ter pelo menos 12 caracteres."
    return None


@app.route("/login", methods=["GET", "POST"])
@limiter.limit(
    "5 per minute",
    exempt_when=lambda: request.method == "GET",
    deduct_when=lambda resposta: resposta.status_code != 302,
)
def login():
    if request.method == "POST":
        nome_usuario = request.form.get("usuario", "").strip()
        senha = request.form.get("senha", "")
        with conectar_banco() as conexao:
            usuario = conexao.execute("""
                SELECT id, nome_usuario, senha_hash, perfil, senha_deve_alterar
                FROM usuarios
                WHERE nome_usuario = ? COLLATE NOCASE AND ativo = 1
            """, (nome_usuario,)).fetchone()

        if usuario and check_password_hash(usuario["senha_hash"], senha):
            session.clear()
            session["usuario_id"] = usuario["id"]
            session["nome_usuario"] = usuario["nome_usuario"]
            session["perfil"] = usuario["perfil"]
            session["senha_deve_alterar"] = bool(usuario["senha_deve_alterar"])
            return redirect(url_for("inicio"))
        flash("Usuário ou senha incorretos.")

    return render_template("login.html")


@app.route("/recuperar-senha")
def recuperar_senha():
    return render_template("recuperar_senha.html")


@app.route("/criar-conta", methods=["GET", "POST"])
@limiter.limit("3 per hour", exempt_when=lambda: request.method == "GET")
def criar_conta():
    if request.method == "POST":
        nome_usuario = request.form.get("usuario", "").strip()
        senha = request.form.get("senha", "")
        erro = validar_credenciais_usuario(nome_usuario, senha)
        if erro:
            flash(erro)
            return redirect(url_for("criar_conta"))

        try:
            with conectar_banco() as conexao:
                conexao.execute("""
                    INSERT INTO usuarios (nome_usuario, senha_hash, perfil)
                    VALUES (?, ?, 'user')
                """, (nome_usuario, generate_password_hash(senha)))
        except sqlite3.IntegrityError:
            flash("Esse nome de usuário já está em uso.")
            return redirect(url_for("criar_conta"))

        flash("Conta criada como usuário. Entre com suas credenciais.")
        return redirect(url_for("login"))

    return render_template("criar_conta.html")


@app.route("/minha-conta/alterar-senha", methods=["GET", "POST"])
def alterar_senha_propria():
    if request.method == "POST":
        senha_atual = request.form.get("senha_atual", "")
        nova_senha = request.form.get("nova_senha", "")
        confirmar_senha = request.form.get("confirmar_senha", "")
        usuario_id = session["usuario_id"]

        with conectar_banco() as conexao:
            usuario = conexao.execute(
                "SELECT nome_usuario, senha_hash FROM usuarios WHERE id = ?",
                (usuario_id,),
            ).fetchone()
            if not check_password_hash(usuario["senha_hash"], senha_atual):
                flash("A senha atual está incorreta.")
                return redirect(url_for("alterar_senha_propria"))
            if len(nova_senha) < 12:
                flash("A nova senha deve ter pelo menos 12 caracteres.")
                return redirect(url_for("alterar_senha_propria"))
            if nova_senha != confirmar_senha:
                flash("A confirmação da nova senha não confere.")
                return redirect(url_for("alterar_senha_propria"))
            if check_password_hash(usuario["senha_hash"], nova_senha):
                flash("Escolha uma senha diferente da atual.")
                return redirect(url_for("alterar_senha_propria"))

            conexao.execute("""
                UPDATE usuarios
                SET senha_hash = ?, senha_deve_alterar = 0
                WHERE id = ?
            """, (generate_password_hash(nova_senha), usuario_id))
            registrar_acao(
                conexao, "Senha alterada",
                f"Senha da conta {usuario['nome_usuario']} alterada."
            )

        session["senha_deve_alterar"] = False
        flash("Senha atualizada com sucesso.")
        return redirect(url_for("minha_conta"))

    return render_template("alterar_senha.html")


@app.route("/sair", methods=["POST"])
def sair():
    session.clear()
    return redirect(url_for("login"))


@app.route("/minha-conta", methods=["GET", "POST"])
def minha_conta():
    usuario_id = session["usuario_id"]
    if request.method == "POST":
        tema = request.form.get("tema", "system")
        cor_botao = request.form.get("cor_botao", "").strip().lower()
        if tema not in {"system", "light", "dark"}:
            flash("Selecione um modo de aparência válido.")
            return redirect(url_for("minha_conta"))
        if not re.fullmatch(r"#[0-9a-f]{6}", cor_botao):
            flash("Escolha uma cor válida para os botões.")
            return redirect(url_for("minha_conta"))

        with conectar_banco() as conexao:
            conexao.execute(
                "UPDATE usuarios SET tema = ?, cor_botao = ? WHERE id = ?",
                (tema, cor_botao, usuario_id)
            )
        session["tema_usuario"] = tema
        session["cor_botao"] = cor_botao
        session["cor_botao_texto"] = cor_texto_botao(cor_botao)
        flash("Preferências salvas.")
        return redirect(url_for("minha_conta"))

    return render_template("minha_conta.html")


@app.route("/minha-conta/tema", methods=["POST"])
def alterar_tema_usuario():
    tema = request.form.get("tema", "")
    if tema not in {"light", "dark", "system"}:
        abort(400)
    with conectar_banco() as conexao:
        conexao.execute(
            "UPDATE usuarios SET tema = ? WHERE id = ?",
            (tema, session["usuario_id"])
        )
    session["tema_usuario"] = tema
    return "", 204

@app.route("/usuarios")
def usuarios():
    with conectar_banco() as conexao:
        lista_usuarios = conexao.execute("""
            SELECT id, nome_usuario, perfil, ativo, criado_em
            FROM usuarios ORDER BY nome_usuario COLLATE NOCASE
        """).fetchall()
    return render_template("usuarios.html", usuarios=lista_usuarios)


@app.route("/usuarios/criar", methods=["POST"])
def criar_usuario():
    nome_usuario = request.form.get("usuario", "").strip()
    senha = request.form.get("senha", "")
    perfil = request.form.get("perfil", "user")

    erro = validar_credenciais_usuario(nome_usuario, senha)
    if erro:
        flash(erro)
        return redirect(url_for("usuarios"))
    if perfil not in {"admin", "user"}:
        flash("Selecione um perfil válido.")
        return redirect(url_for("usuarios"))

    try:
        with conectar_banco() as conexao:
            conexao.execute("""
                INSERT INTO usuarios (nome_usuario, senha_hash, perfil)
                VALUES (?, ?, ?)
            """, (nome_usuario, generate_password_hash(senha), perfil))
            registrar_acao(
                conexao, "Usuário criado",
                f"Conta {nome_usuario} criada com perfil {perfil}."
            )
    except sqlite3.IntegrityError:
        flash("Esse nome de usuário já está em uso.")
        return redirect(url_for("usuarios"))

    flash(f"Usuário {nome_usuario} criado como {perfil}.")
    return redirect(url_for("usuarios"))


@app.route("/usuarios/<int:usuario_id>/perfil", methods=["POST"])
def alterar_perfil_usuario(usuario_id):
    perfil = request.form.get("perfil", "")
    if perfil not in {"admin", "user"}:
        flash("Selecione um perfil válido.")
        return redirect(url_for("usuarios"))
    if usuario_id == session.get("usuario_id") and perfil != "admin":
        flash("Você não pode remover seu próprio perfil de administrador.")
        return redirect(url_for("usuarios"))

    with conectar_banco() as conexao:
        alvo = conexao.execute(
            "SELECT id, nome_usuario, perfil FROM usuarios WHERE id = ?", (usuario_id,)
        ).fetchone()
        if alvo is None:
            flash("Usuário não encontrado.")
            return redirect(url_for("usuarios"))
        if alvo["perfil"] == "admin" and perfil != "admin":
            admins = conexao.execute(
                "SELECT COUNT(*) AS total FROM usuarios WHERE perfil = 'admin' AND ativo = 1"
            ).fetchone()["total"]
            if admins <= 1:
                flash("O sistema precisa manter ao menos um administrador ativo.")
                return redirect(url_for("usuarios"))
        conexao.execute(
            "UPDATE usuarios SET perfil = ? WHERE id = ?", (perfil, usuario_id)
        )
        if alvo["perfil"] != perfil:
            registrar_acao(
                conexao, "Perfil alterado",
                f"Perfil de {alvo['nome_usuario']} alterado de {alvo['perfil']} para {perfil}."
            )

    flash(f"Perfil de {alvo['nome_usuario']} atualizado.")
    return redirect(url_for("usuarios"))


@app.route("/usuarios/<int:usuario_id>/excluir", methods=["POST"])
def excluir_usuario(usuario_id):
    if usuario_id == session.get("usuario_id"):
        flash("Você não pode excluir sua própria conta enquanto estiver conectado.")
        return redirect(url_for("usuarios"))

    with conectar_banco() as conexao:
        alvo = conexao.execute(
            "SELECT id, nome_usuario, perfil FROM usuarios WHERE id = ?", (usuario_id,)
        ).fetchone()
        if alvo is None:
            flash("Usuário não encontrado.")
            return redirect(url_for("usuarios"))
        if alvo["perfil"] == "admin":
            admins = conexao.execute(
                "SELECT COUNT(*) AS total FROM usuarios WHERE perfil = 'admin' AND ativo = 1"
            ).fetchone()["total"]
            if admins <= 1:
                flash("O sistema precisa manter ao menos um administrador ativo.")
                return redirect(url_for("usuarios"))
        conexao.execute("DELETE FROM usuarios WHERE id = ?", (usuario_id,))
        registrar_acao(
            conexao, "Usuário removido",
            f"Conta {alvo['nome_usuario']} removida."
        )

    flash(f"Usuário {alvo['nome_usuario']} removido.")
    return redirect(url_for("usuarios"))


@app.route("/usuarios/<int:usuario_id>/senha", methods=["POST"])
def redefinir_senha_usuario(usuario_id):
    senha_temporaria = request.form.get("senha_temporaria", "")
    if len(senha_temporaria) < 12:
        flash("A senha temporária deve ter pelo menos 12 caracteres.")
        return redirect(url_for("usuarios"))

    with conectar_banco() as conexao:
        usuario = conexao.execute(
            "SELECT nome_usuario FROM usuarios WHERE id = ?", (usuario_id,)
        ).fetchone()
        if usuario is None:
            flash("Usuário não encontrado.")
            return redirect(url_for("usuarios"))

        conexao.execute("""
            UPDATE usuarios
            SET senha_hash = ?, senha_deve_alterar = 1
            WHERE id = ?
        """, (generate_password_hash(senha_temporaria), usuario_id))
        registrar_acao(
            conexao, "Senha redefinida",
            f"Senha temporária definida para {usuario['nome_usuario']}; "
            "a conta deverá alterá-la no próximo acesso."
        )

    flash(
        f"Senha temporária definida para {usuario['nome_usuario']}. "
        "Entregue-a à pessoa por um canal seguro; ela terá que alterá-la ao entrar."
    )
    return redirect(url_for("usuarios"))


def salvar_tipo(conexao, tipo):
    tipo = tipo.strip().upper()

    if tipo:

        conexao.execute("""
            INSERT OR IGNORE INTO tipos_bens (nome)
            VALUES (?)
        """, (tipo,))

    return tipo


def salvar_descricao(conexao, descricao):
    descricao = descricao.strip()

    if descricao:
        conexao.execute("""
            INSERT OR IGNORE INTO descricoes_bens (nome)
            VALUES (?)
        """, (descricao,))

    return descricao


def numero_sala_valido(numero):
    return (
        numero.isascii() and numero.isdigit()
        and len(numero) <= 12 and int(numero) > 0
    )


def validar_dados_bem(patrimonio, descricao, numero_sala, tipo, estado):
    if not patrimonio or not descricao or not numero_sala or not tipo:
        return "Preencha todos os campos, incluindo o tipo."
    if not numero_sala_valido(numero_sala):
        return "Informe um número de sala válido."
    if (
        len(patrimonio) > 120 or len(descricao) > 1000
        or len(tipo) > 120 or len(estado) > 120
    ):
        return "Confira os limites de caracteres dos dados do bem."
    return None


def normalizar_cabecalho(valor):
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(letra for letra in texto if not unicodedata.combining(letra))
    return re.sub(r"[^a-z0-9]", "", texto.lower())


ALIAS_COLUNAS = {
    "patrimonio": {
        "patrimonio", "npatrimonio", "nrpatrimonio", "numeropatrimonio",
        "numerodopatrimonio", "numpatrimonio", "codigopatrimonio",
        "tombamento", "plaqueta", "numerotombo", "codigo", "cod"
    },
    "descricao": {"descricao", "descricaobem", "descricaodobem", "descricaoadicional", "nomebem", "especificacao", "item"},
    "tipo": {"tipo", "tipobem", "categoria", "classe", "descricaoativofixo"},
    "estado": {"estado", "condicao", "estadoconservacao"},
    "etiquetado": {"etiquetado", "possuietiqueta", "temetiqueta", "etiqueta"},
}


def mapa_colunas(cabecalho):
    mapa = {}
    for indice, valor in enumerate(cabecalho):
        nome = normalizar_cabecalho(valor)
        for campo, aliases in ALIAS_COLUNAS.items():
            if nome in aliases and campo not in mapa:
                mapa[campo] = indice
    if "patrimonio" in mapa and "descricao" in mapa:
        return mapa
    return None


def texto_celula(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def extrair_registros(tabelas):
    registros = []
    encontrou_cabecalho = False

    for linhas in tabelas:
        mapa = None
        for linha in linhas:
            linha = list(linha or [])
            novo_mapa = mapa_colunas(linha)
            if novo_mapa:
                mapa = novo_mapa
                encontrou_cabecalho = True
                continue
            if not mapa:
                continue

            def valor(campo):
                indice = mapa.get(campo)
                return texto_celula(linha[indice]) if indice is not None and indice < len(linha) else ""

            patrimonio = valor("patrimonio")
            descricao = valor("descricao")
            if patrimonio or descricao or valor("tipo") or valor("etiquetado") or valor("estado"):
                registros.append({
                    "patrimonio": patrimonio,
                    "descricao": descricao,
                    "tipo": valor("tipo").upper() or "NÃO INFORMADO",
                    "estado": valor("estado"),
                    "etiquetado": valor("etiquetado").lower() in {
                        "sim", "s", "yes", "y", "1", "x", "true", "etiquetado"
                    },
                })

    return registros, encontrou_cabecalho


def ler_excel(conteudo, extensao):
    tabelas = []
    if extensao == ".xlsx":
        from openpyxl import load_workbook
        arquivo = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
        try:
            tabelas = [list(planilha.iter_rows(values_only=True)) for planilha in arquivo.worksheets]
        finally:
            arquivo.close()
    else:
        import xlrd
        arquivo = xlrd.open_workbook(file_contents=conteudo, on_demand=True)
        try:
            tabelas = [
                [planilha.row_values(indice) for indice in range(planilha.nrows)]
                for planilha in arquivo.sheets()
            ]
        finally:
            arquivo.release_resources()
    return extrair_registros(tabelas)


def ler_pdf(conteudo):
    import pdfplumber
    tabelas = []
    with pdfplumber.open(BytesIO(conteudo)) as arquivo:
        for pagina in arquivo.pages:
            tabelas_pagina = pagina.extract_tables() or []
            if tabelas_pagina:
                tabelas.extend(tabelas_pagina)
                continue

            texto = pagina.extract_text() or ""
            linhas = [
                re.split(r"\t+|[|;]+|\s{2,}", linha.strip())
                for linha in texto.splitlines()
                if linha.strip()
            ]
            tabelas.append(linhas)
    return extrair_registros(tabelas)



@app.route("/")
def inicio():
    with conectar_banco() as conexao:
        salas = conexao.execute("""
            SELECT salas.*, COUNT(bens.id) AS total_bens
            FROM salas
            LEFT JOIN bens ON bens.sala_id = salas.id
            GROUP BY salas.id
            ORDER BY CAST(salas.numero AS INTEGER), salas.numero
        """).fetchall()

    return render_template("index.html", salas=salas)


@app.route("/api/historico")
def api_historico():
    with conectar_banco() as conexao:
        registros = conexao.execute("""
            SELECT usuario, acao, detalhe, criado_em
            FROM historico_acoes
            ORDER BY criado_em DESC, id DESC
            LIMIT 50
        """).fetchall()
    resposta = jsonify([
        {
            "usuario": registro["usuario"],
            "acao": registro["acao"],
            "detalhe": registro["detalhe"],
            "criado_em": registro["criado_em"]
        }
        for registro in registros
    ])
    resposta.headers["Cache-Control"] = "no-store"
    return resposta


@app.route("/cadastrar-sala", methods=["POST"])
def cadastrar_sala():
    numero = request.form.get("numero", "").strip()
    nome = request.form.get("nome", "").strip()

    if not numero_sala_valido(numero) or not nome:
        flash("Informe um número válido e o nome da sala.")
        return redirect(url_for("inicio"))

    try:
        with conectar_banco() as conexao:
            conexao.execute("""
                INSERT INTO salas (numero, nome, token_publico)
                VALUES (?, ?, ?)
            """, (numero, nome, secrets.token_urlsafe(24)))
            registrar_acao(
                conexao, "Sala cadastrada",
                f"Sala {numero} - {nome} cadastrada."
            )

        flash("Sala cadastrada com sucesso.")

    except sqlite3.IntegrityError:
        flash("Esse número ou nome de sala já está cadastrado.")

    return redirect(url_for("inicio"))


@app.route("/abrir-sala", methods=["POST"])
def buscar_sala():
    numero = request.form.get("numero_sala", "").strip()

    with conectar_banco() as conexao:
        sala = conexao.execute("""
            SELECT id FROM salas WHERE numero = ?
        """, (numero,)).fetchone()

    if sala is None:
        flash("Sala não encontrada.")
        return redirect(url_for("inicio"))

    return redirect(url_for("abrir_sala", sala_id=sala["id"]))


@app.route("/excluir-sala/<int:sala_id>", methods=["POST"])
def excluir_sala(sala_id):
    with conectar_banco() as conexao:
        sala = conexao.execute("""
            SELECT id, numero, nome FROM salas WHERE id = ?
        """, (sala_id,)).fetchone()

        if sala is None:
            flash("Sala não encontrada.")
            return redirect(url_for("inicio"))

        total_bens = conexao.execute("""
            SELECT COUNT(*) AS total
            FROM bens WHERE sala_id = ?
        """, (sala_id,)).fetchone()["total"]

        if total_bens > 0:
            flash("Não é possível apagar uma sala que contém bens.")
            return redirect(url_for("inicio"))

        conexao.execute(
            "DELETE FROM salas WHERE id = ?", (sala_id,)
        )
        registrar_acao(
            conexao, "Sala removida",
            f"Sala {sala['numero']} - {sala['nome']} removida."
        )

    flash(f"Sala {sala['numero']} - {sala['nome']} apagada.")
    return redirect(url_for("inicio"))


@app.route("/sala/<int:sala_id>")
def abrir_sala(sala_id):
    with conectar_banco() as conexao:
        sala = conexao.execute("""
            SELECT * FROM salas WHERE id = ?
        """, (sala_id,)).fetchone()

        if sala is None:
            return "Sala não encontrada.", 404

        bens = conexao.execute("""
            SELECT * FROM bens
            WHERE sala_id = ?
            ORDER BY patrimonio
        """, (sala_id,)).fetchall()

    return render_template("sala.html", sala=sala, bens=bens)


@app.route("/espelho/<token_publico>")
def espelho_sala(token_publico):
    with conectar_banco() as conexao:
        sala = conexao.execute(
            "SELECT id, numero, nome FROM salas WHERE token_publico = ?", (token_publico,)
        ).fetchone()
        if sala is None:
            return "Sala não encontrada.", 404
        bens = conexao.execute("""
            SELECT patrimonio, descricao, tipo, estado, etiquetado
            FROM bens WHERE sala_id = ? ORDER BY patrimonio
        """, (sala["id"],)).fetchall()
    resposta = app.make_response(render_template("espelho.html", sala=sala, bens=bens))
    resposta.headers["Cache-Control"] = "no-store"
    resposta.headers["X-Robots-Tag"] = "noindex, nofollow"
    resposta.headers["Referrer-Policy"] = "no-referrer"
    return resposta


@app.route("/sala/<int:sala_id>/bem/<int:bem_id>/alternar-etiqueta", methods=["POST"])
def alternar_etiqueta(sala_id, bem_id):
    with conectar_banco() as conexao:
        bem = conexao.execute("""
            SELECT patrimonio, etiquetado FROM bens
            WHERE id = ? AND sala_id = ?
        """, (bem_id, sala_id)).fetchone()
        if bem is None:
            return "Bem não encontrado nesta sala.", 404
        cursor = conexao.execute("""
            UPDATE bens
            SET etiquetado = CASE WHEN etiquetado = 1 THEN 0 ELSE 1 END
            WHERE id = ? AND sala_id = ?
        """, (bem_id, sala_id))
        etiquetado = conexao.execute(
            "SELECT etiquetado FROM bens WHERE id = ? AND sala_id = ?",
            (bem_id, sala_id)
        ).fetchone()["etiquetado"]
        if cursor.rowcount and bem["etiquetado"] != etiquetado:
            registrar_acao(
                conexao, "Etiqueta alterada",
                f"Patrimônio {bem['patrimonio']} marcado como "
                f"{'etiquetado' if etiquetado else 'não etiquetado'}."
            )

    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return jsonify(etiquetado=etiquetado)
    return redirect(url_for("abrir_sala", sala_id=sala_id))


@app.route("/sala/<int:sala_id>/etiquetas-lote", methods=["POST"])
def atualizar_etiquetas_lote(sala_id):
    try:
        bem_ids = sorted({int(valor) for valor in request.form.getlist("bem_ids")})
    except ValueError:
        return "Seleção de bens inválida.", 400

    etiquetado = request.form.get("etiquetado", "")
    if not bem_ids or etiquetado not in {"0", "1"}:
        mensagem = "Selecione ao menos um bem e um status válido."
        if request.headers.get("X-Requested-With") == "XMLHttpRequest":
            return jsonify(erro=mensagem), 400
        flash(mensagem)
        return redirect(url_for("abrir_sala", sala_id=sala_id))

    total_atualizado = 0
    with conectar_banco() as conexao:
        sala = conexao.execute(
            "SELECT id, numero, nome FROM salas WHERE id = ?", (sala_id,)
        ).fetchone()
        if sala is None:
            return "Sala não encontrada.", 404
        for inicio in range(0, len(bem_ids), 500):
            lote = bem_ids[inicio:inicio + 500]
            marcadores = ",".join("?" for _ in lote)
            cursor = conexao.execute(
                f"UPDATE bens SET etiquetado = ? WHERE sala_id = ? AND id IN ({marcadores})",
                (int(etiquetado), sala_id, *lote)
            )
            total_atualizado += cursor.rowcount
        if total_atualizado:
            registrar_acao(
                conexao, "Etiquetas atualizadas",
                f"{total_atualizado} bem(ns) da sala {sala['numero']} - "
                f"{sala['nome']} marcado(s) como "
                f"{'etiquetado' if etiquetado == '1' else 'não etiquetado'}."
            )

    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return jsonify(atualizados=total_atualizado, etiquetado=int(etiquetado))
    flash(f"Status atualizado para {total_atualizado} bem(ns).")
    return redirect(url_for("abrir_sala", sala_id=sala_id))

@app.route("/sala/<int:sala_id>/importar-bens", methods=["POST"])
def importar_bens(sala_id):
    arquivo = request.files.get("arquivo")
    extensao = Path(arquivo.filename or "").suffix.lower() if arquivo else ""

    with conectar_banco() as conexao:
        sala = conexao.execute(
            "SELECT id, numero, nome FROM salas WHERE id = ?", (sala_id,)
        ).fetchone()
        if sala is None:
            return "Sala não encontrada.", 404

        if not arquivo or not arquivo.filename:
            flash("Escolha um arquivo Excel ou PDF para importar.")
            return redirect(url_for("abrir_sala", sala_id=sala_id))

        if extensao not in {".xlsx", ".xls", ".pdf"}:
            flash("Use um arquivo .xlsx, .xls ou .pdf.")
            return redirect(url_for("abrir_sala", sala_id=sala_id))

        try:
            conteudo = arquivo.read(app.config["MAX_CONTENT_LENGTH"] + 1)
            if len(conteudo) > app.config["MAX_CONTENT_LENGTH"]:
                flash("O arquivo deve ter no máximo 8 MB.")
                return redirect(url_for("abrir_sala", sala_id=sala_id))

            if extensao == ".pdf":
                registros, encontrou_cabecalho = ler_pdf(conteudo)
            else:
                registros, encontrou_cabecalho = ler_excel(conteudo, extensao)
        except ModuleNotFoundError:
            flash("Faltam bibliotecas da importação. No terminal, execute: python -m pip install -r requirements.txt")
            return redirect(url_for("abrir_sala", sala_id=sala_id))
        except Exception:
            flash("Não foi possível ler o arquivo. Confira se ele não está protegido ou corrompido.")
            return redirect(url_for("abrir_sala", sala_id=sala_id))

        if not encontrou_cabecalho:
            flash("Não encontrei as colunas Patrimônio e Descrição no arquivo.")
            return redirect(url_for("abrir_sala", sala_id=sala_id))
        if not registros:
            flash("Encontrei os cabeçalhos, mas nenhuma linha válida para importar.")
            return redirect(url_for("abrir_sala", sala_id=sala_id))

        inseridos = 0
        duplicados = 0
        invalidos = 0
        vistos = set()

        for registro in registros:
            patrimonio = registro["patrimonio"].strip()
            descricao = registro["descricao"].strip()
            tipo = registro["tipo"].strip().upper()
            estado = registro["estado"].strip()

            if validar_dados_bem(patrimonio, descricao, "1", tipo, estado):
                invalidos += 1
                continue
            if patrimonio.casefold() in vistos:
                duplicados += 1
                continue
            vistos.add(patrimonio.casefold())

            try:
                conexao.execute("""
                    INSERT INTO bens (patrimonio, descricao, sala_id, etiquetado, tipo, estado)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (
                    patrimonio, descricao, sala_id,
                    1 if registro["etiquetado"] else 0,
                    tipo, estado
                ))
                salvar_tipo(conexao, tipo)
                salvar_descricao(conexao, descricao)
                inseridos += 1
            except sqlite3.IntegrityError:
                duplicados += 1

        if inseridos:
            registrar_acao(
                conexao, "Bens importados",
                f"{inseridos} bem(ns) importado(s) para a sala "
                f"{sala['numero']} - {sala['nome']}."
            )

    flash(
        f"Importação concluída: {inseridos} item(ns) cadastrado(s), "
        f"{duplicados} duplicado(s) ignorado(s) e {invalidos} linha(s) inválida(s)."
    )
    return redirect(url_for("abrir_sala", sala_id=sala_id))


@app.route("/sala/<int:sala_id>/conferencia")
def conferencia_sala(sala_id):
    with conectar_banco() as conexao:
        sala = conexao.execute("""
            SELECT * FROM salas WHERE id = ?
        """, (sala_id,)).fetchone()

        if sala is None:
            return "Sala não encontrada.", 404

        bens = conexao.execute("""
            SELECT * FROM bens
            WHERE sala_id = ?
            ORDER BY patrimonio
        """, (sala_id,)).fetchall()

    return render_template("conferencia.html", sala=sala, bens=bens)


@app.route("/sala/<int:sala_id>/conferencia/<int:bem_id>", methods=["POST"])
def marcar_conferencia(sala_id, bem_id):
    conferido = 1 if request.form.get("conferido") == "1" else 0

    with conectar_banco() as conexao:
        bem = conexao.execute("""
            SELECT patrimonio, conferido FROM bens
            WHERE id = ? AND sala_id = ?
        """, (bem_id, sala_id)).fetchone()
        cursor = conexao.execute("""
            UPDATE bens
            SET conferido = ?
            WHERE id = ? AND sala_id = ?
        """, (conferido, bem_id, sala_id))
        if cursor.rowcount and bem and bem["conferido"] != conferido:
            registrar_acao(
                conexao, "Conferência atualizada",
                f"Patrimônio {bem['patrimonio']} "
                f"{'marcado como conferido' if conferido else 'removido da conferência'}."
            )

    if not cursor.rowcount:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest":
            return jsonify(erro="Bem não encontrado nesta sala."), 404
        flash("Bem não encontrado nesta sala.")
    elif request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return jsonify(conferido=bool(conferido))
    else:
        flash("Item marcado como conferido." if conferido else "Marcação removida.")

    return redirect(url_for("conferencia_sala", sala_id=sala_id))

@app.route("/pesquisar-bens")
def pesquisar_bens():
    termo = request.args.get("q", "").strip()
    bem_id = request.args.get("bem_id", type=int)
    bem_selecionado = None

    with conectar_banco() as conexao:
        salas = conexao.execute("""
            SELECT id, numero, nome
            FROM salas
            ORDER BY CAST(numero AS INTEGER), numero
        """).fetchall()

        tipos = conexao.execute("""
            SELECT nome FROM tipos_bens
            ORDER BY nome
        """).fetchall()

        descricoes = conexao.execute("""
            SELECT nome FROM descricoes_bens
            ORDER BY nome COLLATE NOCASE
        """).fetchall()

        bens = []

        if termo:
            busca = f"%{termo}%"
            bens = conexao.execute("""
                SELECT bens.*,
                       salas.numero AS numero_sala,
                       salas.nome AS nome_sala
                FROM bens
                JOIN salas ON salas.id = bens.sala_id
                WHERE bens.patrimonio LIKE ?
                   OR bens.descricao LIKE ?
                   OR bens.tipo LIKE ?
                   OR salas.numero LIKE ?
                   OR salas.nome LIKE ?
                ORDER BY bens.patrimonio
            """, (busca, busca, busca, busca, busca)).fetchall()

            if bem_id is not None:
                bem_selecionado = next(
                    (bem for bem in bens if bem["id"] == bem_id), None
                )
            elif len(bens) == 1:
                bem_selecionado = bens[0]

    return render_template(
        "pesquisar_bens.html",
        bens=bens,
        salas=salas,
        tipos=tipos,
        descricoes=descricoes,
        termo=termo,
        bem_selecionado=bem_selecionado
    )


@app.route("/cadastrar-bem", methods=["POST"])
def cadastrar_bem():
    patrimonio = request.form.get("patrimonio", "").strip()
    descricao = request.form.get("descricao", "").strip()
    numero_sala = request.form.get("numero_sala", "").strip()
    tipo = request.form.get("tipo", "").strip()
    estado = request.form.get("estado", "").strip()
    etiquetado = 1 if request.form.get("etiquetado") == "1" else 0

    erro = validar_dados_bem(patrimonio, descricao, numero_sala, tipo, estado)
    if erro:
        flash(erro)
        return redirect(url_for("pesquisar_bens"))

    with conectar_banco() as conexao:
        sala = conexao.execute("""
            SELECT id FROM salas WHERE numero = ?
        """, (numero_sala,)).fetchone()

        if sala is None:
            flash("A sala informada não existe.")
            return redirect(url_for("pesquisar_bens"))

        tipo = tipo.strip().upper()
        descricao = descricao.strip()
        cadastrado = False

        try:
            conexao.execute("""
                INSERT INTO bens
                    (patrimonio, descricao, sala_id, etiquetado, tipo, estado)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (
                patrimonio, descricao, sala["id"], etiquetado, tipo, estado
            ))

            flash("Bem cadastrado com sucesso.")
            cadastrado = True
            salvar_tipo(conexao, tipo)
            salvar_descricao(conexao, descricao)
            registrar_acao(
                conexao, "Bem cadastrado",
                f"Patrimônio {patrimonio} cadastrado na sala {numero_sala}."
            )

        except sqlite3.IntegrityError:
            flash("Esse número patrimonial já está cadastrado.")

    if cadastrado:
        return redirect(url_for("pesquisar_bens"))

    return redirect(url_for("pesquisar_bens", q=patrimonio))


@app.route("/editar-bem/<int:bem_id>", methods=["POST"])
def editar_bem(bem_id):
    patrimonio = request.form.get("patrimonio", "").strip()
    descricao = request.form.get("descricao", "").strip()
    numero_sala = request.form.get("numero_sala", "").strip()
    tipo = request.form.get("tipo", "").strip()
    estado = request.form.get("estado", "").strip()
    etiquetado = 1 if request.form.get("etiquetado") == "1" else 0
    termo = request.form.get("termo", "").strip()

    erro = validar_dados_bem(patrimonio, descricao, numero_sala, tipo, estado)
    if erro:
        flash(erro)
        return redirect(url_for("pesquisar_bens", q=termo))

    with conectar_banco() as conexao:
        sala = conexao.execute("""
            SELECT id FROM salas WHERE numero = ?
        """, (numero_sala,)).fetchone()

        if sala is None:
            flash("A sala informada não existe.")
            return redirect(url_for("pesquisar_bens", q=termo))

        tipo = tipo.strip().upper()
        descricao = descricao.strip()
        atualizado = False
        bem_existente = conexao.execute(
            "SELECT patrimonio FROM bens WHERE id = ?", (bem_id,)
        ).fetchone()

        try:
            cursor = conexao.execute("""
                UPDATE bens
                SET patrimonio = ?,
                    descricao = ?,
                    sala_id = ?,
                    etiquetado = ?,
                    tipo = ?,
                    estado = ?
                WHERE id = ?
            """, (
                patrimonio, descricao, sala["id"],
                etiquetado, tipo, estado, bem_id
            ))

            if cursor.rowcount:
                flash("Bem atualizado com sucesso.")
                atualizado = True
                salvar_tipo(conexao, tipo)
                salvar_descricao(conexao, descricao)
                registrar_acao(
                    conexao, "Bem atualizado",
                    f"Patrimônio {bem_existente['patrimonio']} atualizado para {patrimonio}."
                )
            else:
                flash("Bem não encontrado.")

        except sqlite3.IntegrityError:
            flash("Esse número patrimonial já pertence a outro bem.")

    if atualizado:
        return redirect(url_for("pesquisar_bens"))

    return redirect(
        url_for("pesquisar_bens", q=termo or patrimonio, bem_id=bem_id)
    )


@app.route("/excluir-bem/<int:bem_id>", methods=["POST"])
def excluir_bem(bem_id):
    termo = request.form.get("termo", "").strip()

    with conectar_banco() as conexao:
        bem = conexao.execute("""
            SELECT id, patrimonio FROM bens WHERE id = ?
        """, (bem_id,)).fetchone()

        if bem is None:
            flash("Bem não encontrado.")
            return redirect(url_for("pesquisar_bens", q=termo))

        conexao.execute(
            "DELETE FROM bens WHERE id = ?", (bem_id,)
        )
        registrar_acao(
            conexao, "Bem removido",
            f"Patrimônio {bem['patrimonio']} removido."
        )

    flash("Bem excluído com sucesso.")
    return redirect(url_for("pesquisar_bens", q=termo))


@app.route("/sala/<int:sala_id>/qrcode.svg")
def gerar_qr(sala_id):
    with conectar_banco() as conexao:
        sala = conexao.execute("""
            SELECT token_publico FROM salas WHERE id = ?
        """, (sala_id,)).fetchone()

    if sala is None:
        return "Sala não encontrada.", 404

    base_url = app.config["PUBLIC_BASE_URL"] or request.url_root.rstrip("/")
    link_sala = base_url + url_for("espelho_sala", token_publico=sala["token_publico"])

    factory = qrcode.image.svg.SvgPathImage
    imagem = qrcode.make(link_sala, image_factory=factory, box_size=8, border=4)
    arquivo = BytesIO(imagem.to_string())
    arquivo.seek(0)

    return send_file(
        arquivo,
        mimetype="image/svg+xml",
        max_age=0
    )


if __name__ == "__main__":
    criar_banco()

    print("\nSistema de Imobilizado QR iniciado.")
    print(f"Computador: http://127.0.0.1:{PORTA}")
    endereco = app.config["PUBLIC_BASE_URL"] or f"http://127.0.0.1:{PORTA}"
    print(f"Endereço configurado: {endereco}\n")

    app.run(host="0.0.0.0", port=PORTA, debug=False)
